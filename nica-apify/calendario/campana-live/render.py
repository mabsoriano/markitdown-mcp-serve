import importlib.util, html, re, sys, os
here = os.path.dirname(os.path.abspath(__file__))
s = importlib.util.spec_from_file_location('c', os.path.join(here, 'contenido.py')); m = importlib.util.module_from_spec(s); s.loader.exec_module(m)
E = lambda t: re.sub(r'\[(.+?)\]', r'<mark>[\1]</mark>', html.escape(t or ''))
CHIP = {'A': ('Tú a cámara · decreto o ritual', 'a'), 'B': ('Tú a cámara · rompe una creencia', 'b'), 'Y': ('Yapping · platicando', 'y'), 'R': ('Automático · tú no grabas', 'r')}
uid = [0]
def card(p):
    uid[0] += 1; i = uid[0]
    lab, cls = CHIP[p['n']]
    cta = p['cta']
    ctacls = 'live' if cta.startswith('LIVE') else ('venta' if 'Robóticas' in cta else 'none')
    head = f'<summary><span class="hora">{p["hora"]}</span><span class="tit">{E(p["titulo"])}</span><span class="chips"><span class="chip {cls}">{lab}</span><span class="chip {ctacls}">CTA: {html.escape(cta)}</span></span></summary>'
    body = [f'<p class="hook"><span class="lbl">Primera frase</span>“{E(p["hook"])}”</p>', f'<p class="quien"><span class="lbl">Quién lo hace</span>{E(p["quien"])}</p>']
    if p['n'] == 'Y':
        body.append('<div class="blk"><span class="lbl">Puntos para platicar (sin guion palabra por palabra)</span><ol class="pts">' + ''.join(f'<li>{E(x)}</li>' for x in p['puntos']) + '</ol></div>')
        body.append(f'<p><span class="lbl">Cierra con esta pregunta</span>{E(p["pregunta"])}</p>')
        body.append('<p class="nota">Sin venta y sin link: solo conexión. Un solo corte, cerca de la cámara (coche, cocina, caminando).</p>')
    else:
        body.append('<div class="blk"><span class="lbl">Guion por tiempos</span><ol class="guion">' + ''.join(f'<li><span class="tc">{tc}</span><span>{E(tx)}</span></li>' for tc, tx in p['guion']) + '</ol></div>')
        body.append(f'<div class="blk"><span class="lbl">Texto en pantalla</span><p>{E(p["pantalla"])}</p></div>')
    body.append(f'<div class="blk"><div class="row"><span class="lbl">Copy para publicar</span><button type="button" class="cp" data-t="c{i}">Copiar copy</button></div><pre id="c{i}">{html.escape(p["copy"])}\n\n{html.escape(p["hashtags"])}</pre></div>')
    return f'<details class="reel {cls}">{head}<div class="body">{"".join(body)}</div></details>'
def semana(sid, titulo, sub, dias):
    graba = sum(1 for d in dias for p in d[3] if p['n'] != 'R')
    out = [f'<section class="semana" id="{sid}"><h2>{titulo}</h2><p class="sub">{sub} · Tú grabas {graba} piezas; {sum(1 for d in dias for p in d[3] if p["n"]=="R")} se generan solas.</p>']
    for fecha, dia, tema, piezas in dias:
        out.append(f'<div class="dia"><h3>{dia} <span class="tema">{tema}</span></h3>' + ''.join(card(p) for p in sorted(piezas, key=lambda p: p['hora'])) + '</div>')
    return ''.join(out) + '</section>'
A = m.ANUNCIOS
def ads():
    o = ['<section class="semana" id="anuncios"><h2>Anuncios · Facebook e Instagram</h2><p class="sub">Propuesta: nada se activa sin tu OK. Mismos anuncios en las 2 fases; solo cambia la fecha del LIVE.</p>']
    o.append('<div class="grid2">' + ''.join(f'<div class="fase"><b>{html.escape(f["nombre"])}</b><span>Corre: {html.escape(f["fechas"])}</span><span>LIVE: {html.escape(f["fecha_live"])}</span></div>' for f in A['fases']) + '</div>')
    o.append('<div class="aviso"><strong>Reglas de Meta para que no te rechacen el anuncio</strong><ul>' + ''.join(f'<li>{E(r)}</li>' for r in A['reglas']) + '</ul></div>')
    o.append('<h3>3 videos tuyos (tú a cámara)</h3>')
    for v in A['videos']:
        uid[0] += 1
        o.append(f'<details class="reel a"><summary><span class="hora">{v["id"]}</span><span class="tit">{E(v["nombre"])}</span><span class="chips"><span class="chip a">{html.escape(v["duracion"])}</span></span></summary><div class="body"><ol class="guion">' + ''.join(f'<li><span class="tc">{tc}</span><span>{E(tx)}</span></li>' for tc, tx in v['guion']) + f'</ol><p><span class="lbl">Texto en pantalla</span>{E(v["pantalla"])}</p></div></details>')
    o.append('<h3>2 imágenes (diseños que ya aprobaste)</h3><ul class="lista">' + ''.join(f'<li><b>{i["id"]} · {E(i["nombre"])}</b>: {E(i["detalle"])}</li>' for i in A['imagenes']) + '</ul>')
    o.append('<h3>Textos del anuncio (Meta prueba los 3)</h3>')
    for t in A['textos']:
        uid[0] += 1
        o.append(f'<div class="blk copyblk"><div class="row"><span class="lbl">Texto {A["textos"].index(t)+1}</span><button type="button" class="cp" data-t="c{uid[0]}">Copiar</button></div><pre id="c{uid[0]}">{html.escape(t)}</pre></div>')
    o.append('<p><span class="lbl">Títulos (máx. 40 caracteres)</span>' + ' · '.join(E(t) for t in A['titulos']) + '</p>')
    f = A['formulario']
    o.append(f'<h3>Formulario de registro dentro del anuncio</h3><ul class="lista"><li><b>Nombre del formulario:</b> {E(f["nombre"])}</li><li><b>Introducción:</b> {E(f["intro"])}</li><li><b>Preguntas:</b> {E(" · ".join(f["preguntas"]))}</li><li><b>Pantalla final:</b> {E(f["final"])}</li></ul>')
    o.append(f'<p><span class="lbl">Público</span>{E(A["publico"])}</p><p><span class="lbl">Regla para apagar</span>{E(A["regla"])}</p>')
    o.append('''<table><thead><tr><th>Presupuesto</th><th>Por día</th><th>Total 2 fases (12 días)</th><th>Registros estimados*</th></tr></thead><tbody>
<tr><td>A</td><td>$200</td><td>$2,400</td><td>60 a 160</td></tr><tr class="rec"><td>B · recomendada</td><td>$300</td><td>$3,600</td><td>90 a 240</td></tr><tr><td>C</td><td>$500</td><td>$6,000</td><td>150 a 400</td></tr></tbody></table>
<p class="nota">*Estimación, no promesa: no tienes datos propios de anuncios de registro. Referencia de $15 a $40 MXN por registro y 20% a 35% de asistencia.</p></section>''')
    return ''.join(o)
tpl = open(os.path.join(here, 'plantilla.html')).read()
page = tpl.replace('%%S1%%', semana('semana1', 'Semana 1 · LIVE martes 3 de noviembre', 'Jueves 29 oct al martes 3 nov', m.S1)) \
          .replace('%%S2%%', semana('semana2', 'Semana 2 · LIVE martes 10 de noviembre', 'Miércoles 4 al martes 10 nov · fecha propuesta, falta tu OK', m.S2)) \
          .replace('%%ADS%%', ads())
open(sys.argv[1], 'w').write(page)
# Excel (regla: entrega también en Excel)
import openpyxl
from openpyxl.styles import Font, Alignment
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Calendario'
ws.append(['Fecha', 'Día', 'Hora', 'Tipo', 'Quién', 'TÍTULO HOOK', 'Primera frase', 'Guion / puntos', 'Texto en pantalla', 'Copy', 'Hashtags', 'CTA'])
for fecha, dia, tema, piezas in m.S1 + m.S2:
    for p in sorted(piezas, key=lambda p: p['hora']):
        g = '\n'.join(f'{a} {b}' for a, b in p['guion']) if p['n'] != 'Y' else '\n'.join(p['puntos']) + '\nCierre: ' + p['pregunta']
        ws.append([fecha, dia, p['hora'], p['tipo'], p['quien'], p['titulo'], p['hook'], g, p.get('pantalla', ''), p['copy'], p['hashtags'], p['cta']])
for c in ws[1]: c.font = Font(bold=True)
for col, w in zip('ABCDEFGHIJKL', [11, 18, 7, 22, 22, 34, 34, 70, 40, 40, 30, 14]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws2 = wb.create_sheet('Anuncios')
for v in A['videos']: ws2.append([v['id'], v['nombre'], v['duracion'], '\n'.join(f'{a} {b}' for a, b in v['guion']), v['pantalla']])
for t in A['textos']: ws2.append(['Texto', t])
ws2.append(['Títulos', ' | '.join(A['titulos'])])
for k, v in A['formulario'].items(): ws2.append(['Formulario · ' + k, v if isinstance(v, str) else ' · '.join(v)])
ws2.append(['Público', A['publico']]); ws2.append(['Regla', A['regla']])
for r in A['reglas']: ws2.append(['Regla Meta', r])
wb.save(sys.argv[2])
print('ok', uid[0])
